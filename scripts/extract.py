#!/usr/bin/env python3
"""Extract the four presentation sections of Кегейли.xlsx and the client's whole-project summary.

Writes
  src/data/kegeyli.json  every value with its sheet!cell, Excel formula and notes
  docs/verification.md   the same values laid out like the sheet, to check against Excel

Run: .venv/bin/python scripts/extract.py [path/to/Кегейли.xlsx]
"""
from __future__ import annotations

import datetime as dt
import hashlib
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SOURCE = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "data" / "Кегейли.xlsx"
OUT_JSON = ROOT / "src" / "data" / "kegeyli.json"
OUT_MD = ROOT / "docs" / "verification.md"

# Cost columns, identical on all four sheets. Each facility uses three rows:
# the facility itself, then +1 own funds, then +2 bank credit.
COST_COLS = {"I": "total", "J": "construction", "K": "equipment", "L": "chickens", "M": "feed"}
MONEY_ROWS = {"cost": 0, "own": 1, "bank": 2}

# Presentation order: grandparent -> parent -> broiler, then the separate egg line.
SECTIONS = [
    {
        "id": "praroditel", "sheet": "Прородитель", "title": "Прародитель",
        "subtitle": "Бройлер йўналиши бўйича прародитель галаси",
        "rows": [10, 13, 16], "birds": ("birdsPerBuilding", "birds"),
        "units": {
            "cost": {"unit": "kUSD", "source": "Прородитель!K3"},
            "revenue": {"unit": "kUSD", "note": "No unit in the sheet; read as thousand USD (3 000 000 chicks × $7), confirmed by the user"},
            "output": {"unit": "pcs", "what": "parent-stock chicks per year"},
        },
        "labels": {"D": "Бош (1 бинода)", "E": "Жами бош", "F": "Родитель жўжа, бош/йил", "G": "Тушум, минг $", "L": "Товуқ"},
    },
    {
        "id": "roditel", "sheet": "Родилеь", "title": "Родитель",
        "subtitle": "Бройлер йўналиши бўйича родитель галаси",
        "rows": [10, 13, 16], "birds": ("birdsPerBuilding", "birds"),
        "units": {
            "cost": {"unit": "kUSD", "note": "No unit in the sheet; read as thousand USD like броллер!K3 and Прородитель!K3"},
            "revenue": {"unit": "USD", "source": "Родилеь!G5"},
            "output": {"unit": "pcs", "what": "hatching eggs per year"},
        },
        "labels": {"D": "Бош (1 бинода)", "E": "Жами бош", "F": "Тухум, дона/йил", "G": "Тушум, $", "L": "Товуқ"},
    },
    {
        "id": "broiler", "sheet": "броллер", "title": "Бройлер",
        "subtitle": "Бройлер йўналиши",
        "rows": [10, 13, 16, 19, 22], "birds": ("birdsPerBatch", "birdsPerYear"),
        "units": {
            "cost": {"unit": "kUSD", "source": "броллер!K3"},
            "revenue": {"unit": "USD", "source": "броллер!G5"},
            "output": {"unit": "kg", "what": "live weight per year", "note": "Header says тн but F = E × 2.5 kg, so the values are kg; slides show tonnes"},
        },
        "labels": {"D": "1 боқишда, бош", "E": "Йилига (×6), бош", "F": "Тирик вазн, кг", "G": "Тушум, $", "L": "Жўжа"},
    },
    {
        "id": "nesushka", "sheet": "Несушка", "title": "Тухум",
        "subtitle": "Тухум йўналиши",
        "rows": [10, 13, 16], "birds": ("birdsPerBuilding", "birds"),
        "units": {
            "cost": {"unit": "kUSD", "note": "No unit in the sheet; read as thousand USD like броллер!K3 and Прородитель!K3"},
            "revenue": {"unit": "USD", "source": "Несушка!G5"},
            "output": {"unit": "pcs", "what": "table eggs per year"},
        },
        "labels": {"D": "Бош (1 бинода)", "E": "Жами бош", "F": "Тухум, дона/йил", "G": "Тушум, $", "L": "Товуқ"},
    },
]

# Sums of "birds per building" across facilities: not a real total, never shown.
HIDDEN = {"Прородитель!D7", "Родилеь!D7"}

# The client's summary of the whole project. The cover and the Жами slide show its totals, so they
# match the file: the four sections plus processing, the feed reserve and the generator. Lines are
# found by their text in column B, because rows move (they did when the generator was added).
PROJECT_SHEET = "Жами Лойиҳа Кегели"
PROJECT_COLS = {"D": "total", "E": "construction", "F": "equipment", "G": "chickens", "H": "feed"}
PROJECT_LINES = [  # column B starts with, id, name on the slides; section lines must match their sheets
    ("Прородитель", "praroditel", "Прародитель"),
    ("Родитель", "roditel", "Родитель"),
    ("Броллер", "broiler", "Бройлер"),
    ("Тухум", "nesushka", "Тухум"),
    ("Дастгох", "processing", "Қайта ишлаш"),
    ("Озуқа заҳираси", "feedReserve", "Озуқа заҳираси"),
    ("Генератор", "generator", "Генератор"),
]
PROJECT_FLAGS = [
    {"level": "total", "refs": [f"{PROJECT_SHEET}!D18:H20", "дастгох!F7:H9"],
     "text": "The processing line (Дастгох) is typed in: 5 657 = 700 own + 4 957 bank. The дастгох sheet gives 5 657,5 "
             "(H7 = 4 457,5), and its own F8/F9 totals (400, 800) skip rows.",
     "resolution": "Using the summary's 5 657, shown as «Қайта ишлаш» (the дастгох sheet's own title). The 0,5 doesn't show at млн $."},
    {"level": "minor", "refs": [f"{PROJECT_SHEET}!H29"],
     "text": "Formula has H179 instead of H17 (Родитель, feed, bank).", "resolution": "H17 = 0, so no effect."},
    {"level": "minor", "refs": [f"{PROJECT_SHEET}!F10", f"{PROJECT_SHEET}!H10", f"{PROJECT_SHEET}!H11"],
     "text": "The Несушка line reads Прородитель!K8, M8, M9 instead of Несушка's cells.", "resolution": "All are 0 on both sheets; no effect."},
    {"level": "minor", "refs": [f"{PROJECT_SHEET}!F29"],
     "text": "Adds F24 (the generator's total) instead of F26 (its bank row).", "resolution": "Same value: the generator has no own funds."},
    {"level": "minor", "refs": [f"{PROJECT_SHEET}!E28:H28"],
     "text": "The own-funds totals skip rows that are 0 (feed reserve, generator, some equipment cells).", "resolution": "No effect."},
    {"level": "text", "refs": ["жами лойиха"],
     "text": "An older three-line summary (Бройлер, Прародитель, Қайта ишлаш; жами кредит 22 042,5).",
     "resolution": f"Not used: «{PROJECT_SHEET}» is the current summary."},
]

# Changes the client asked for on top of the sheet. Money moves from one cell to another,
# then every formula on that sheet is recalculated, so all totals and slides follow.
# Example: {"sheet": "Прородитель", "move": 1000, "from": "L12", "to": "L11", "why": "..."}.
# Empty since the 27.09.2026 file: the client made the change in the sheet itself (1 800 moved in L11/L12).
ADJUSTMENTS: list[dict] = []

# Spelling fixes applied to every text shown on the slides. The sheet's own text is
# kept next to it in the JSON as "sheetText".
WORD_FIXES = [
    ("Родилеь", "Родитель"), ("родилеь", "родитель"),
    ("Прородитель", "Прародитель"), ("прородитель", "прародитель"),
    ("Кеегейли", "Кегейли"), ("Трик вазнда", "Тирик вазнда"),
    ("ота на", "ота-она"), ("ота она", "ота-она"), ("доллрдан", "доллардан"),
    ("Дастгох", "Дастгоҳ"), ("хар бир", "ҳар бир"), ("Лойиха", "лойиҳа"),
    ("АКШ", "АҚШ"), ("2,5кг", "2,5 кг"),
]

# Header cells whose sheet text is misleading or too long for a slide.
HEADER_TEXT = {
    "I4": "Лойиҳа қиймати",
    "H4": "Дастгоҳ етказиб берувчи",
    "J4": "Харажатлар таркиби",  # the sheet repeats «Лойиҳа йўналиши» from B4 here
    "G5": "Тушум",
    "K3": "минг АҚШ доллари",
    "броллер!E5": "Йилига 6 маротаба боқилади",
    "броллер!F5": "Тирик вазнда (2,5 кг)",
    "Родилеь!F5": "Йиллик тухум, дона",
    "Несушка!F5": "Йиллик тухум, дона",
    "Прородитель!F5": "Йиллик родитель жўжа ишлаб чиқариш қуввати, бош",
    "Родилеь!L5": "Товуқ қиймати (ҳар бир бош жўжа 7 доллардан)",
    "Родилеь!M5": "Озуқа ва вакцина (ҳар бир бошга ўртача 7 доллардан)",
}
HEADER_CELLS = ["C4", "D4", "E4", "E5", "F5", "G5", "H4", "I4", "J4", "J5", "K5", "L5", "M5", "K3"]

# Facility names: the sheet text (whitespace collapsed) and what the slides show.
# If the sheet text changes, the script warns and falls back to the sheet text.
FARM = {"kind": "farm", "district": "Кегейли тумани"}
FACILITIES = {
    "Прородитель!B10": ("Кегейли тумани 1-фабрика прородитель (бир бино ўлчами 100*18)",
                        {**FARM, "name": "1-фабрика — прародитель", "size": "100 × 18 м"}),
    "Прородитель!B13": ("Инкубатория прородитель қуввати йилиги 3 млн дона",
                        {"kind": "hatchery", "name": "Инкубация цехи", "capacity": "йиллик қуввати 3 млн дона"}),
    "Прородитель!B16": ("Кеегейли тумани 2-фабрика рем молодняк (бир бино ўлчами 90*14,5)",
                        {**FARM, "name": "2-фабрика — рем молодняк", "size": "90 × 14,5 м"}),
    "Родилеь!B10": ("Кегейли тумани 1-фабрика родилеь (рем молодняк) (бир бино ўлчами 100*18)",
                    {**FARM, "name": "1-фабрика — родитель (рем молодняк)", "size": "100 × 18 м"}),
    "Родилеь!B13": ("Кегейли тумани 2-фабрика родитель (бир бино ўлчами 100*18)",
                    {**FARM, "name": "2-фабрика — родитель", "size": "100 × 18 м"}),
    "Родилеь!B16": ("Инкубатория бройлер қуввати 12 млн бош",
                    {"kind": "hatchery", "name": "Инкубация цехи", "capacity": "бройлер қуввати 12 млн бош"}),
    **{
        f"броллер!B{row}": (f"Кегейли тумани {n}-фабрика (бир бино ўлчами 100*18)",
                            {**FARM, "name": f"{n}-фабрика", "size": "100 × 18 м"})
        for n, row in enumerate([10, 13, 16, 19, 22], start=1)
    },
    "Несушка!B10": ("Кегейли тумани 1-фабрика (бир бино ўлчами 110*18)",
                    {**FARM, "name": "1-фабрика", "size": "110 × 18 м"}),
    "Несушка!B13": ("Кегейли тумани 2-фабрика (бир бино ўлчами 110*18)",
                    {**FARM, "name": "2-фабрика", "size": "110 × 18 м"}),
    "Несушка!B16": ("Кегейли тумани 3-фабрика (рем молодняк) (бир бино ўлчами 110*18)",
                    {**FARM, "name": "3-фабрика — рем молодняк", "size": "110 × 18 м"}),
}

SUPPLIERS = {
    "Биг хердсман Хитой": "Биг Хердсман, Хитой",
    "Е-фарминг Хитой": "Е-Фарминг, Хитой",
    "Е-Фарминг Хитой": "Е-Фарминг, Хитой",
    "Гуанжу поултринг Хитой": "Гуанжу Поултринг, Хитой",
    "Гунажу поултринг Хитой": "Гуанжу Поултринг, Хитой",
}

# Findings that need a human decision; the automatic ones (typed-in totals) are added at runtime.
STATIC_FLAGS = [
    {"level": "unit", "refs": ["Родилеь!I7:M18", "Несушка!I7:M18"],
     "text": "No unit on the cost columns. броллер!K3 and Прородитель!K3 say «1000 АКШ доллари»; Родилеь!L13 = 180 000 birds × $7 = 1 260 agrees.",
     "resolution": "Read as thousand USD."},
    {"level": "unit", "refs": ["Прородитель!G7", "Прородитель!G10"],
     "text": "Revenue has no unit; the other sheets give revenue in whole dollars. 3 000 000 parent chicks × $7 = $21 000 000.",
     "resolution": "Read as thousand USD ($21 млн), confirmed by you."},
    {"level": "unit", "refs": ["броллер!F5", "броллер!F7", "броллер!F10:F22"],
     "text": "Header says тн, but F = E × 2.5 kg per bird, so the values are kg. G = F × 1.8 ($ per kg) agrees.",
     "resolution": "Kept as kg in the JSON; slides show tonnes (15 750 т)."},
    {"level": "total", "refs": sorted(HIDDEN),
     "text": "Adds up «birds per building» across facilities; not a real total.",
     "resolution": "Not shown. Slides use E7 (Родитель 300 000, Прародитель 90 000)."},
    {"level": "total", "refs": ["Родилеь!C7", "Прородитель!C7"],
     "text": "Building count includes the incubator.",
     "resolution": "Shown with that note."},
    {"level": "minor", "refs": ["Родилеь!F7", "Несушка!F7", "Прородитель!F7"],
     "text": "Formula skips row 22, which is empty.", "resolution": "No effect."},
    {"level": "minor", "refs": ["броллер!J10"],
     "text": "Formula is =+J11++J12 (double plus).", "resolution": "No effect."},
    {"level": "minor", "refs": ["броллер!D13", "броллер!D16", "броллер!D19", "броллер!D22"],
     "text": "Typed-in 210 000, while D10 is =C10*70000.", "resolution": "Same value; no effect."},
    {"level": "minor", "refs": ["броллер!A23", "броллер!A24"],
     "text": "Numbered 6 and 7 instead of 5.1 and 5.2.", "resolution": "Numbering not used."},
    {"level": "minor", "refs": ["Родилеь!A19:M24", "Несушка!A19:M24", "Прородитель!A19:M24", "броллер!A25:M27", "броллер!B35"],
     "text": "Empty placeholder rows, and a stray «Жами» in броллер!B35.", "resolution": "Skipped."},
    {"level": "text", "refs": ["*!J4"],
     "text": "The cost group header repeats «Лойиҳа йўналиши» from B4.",
     "resolution": "Shown as «Харажатлар таркиби»."},
    {"level": "text", "refs": ["B10:B22"],
     "text": "Building sizes (100*18, 90*14,5, 110*18) have no unit.",
     "resolution": "Shown in metres (100 × 18 м)."},
    {"level": "text", "refs": ["sheet names", "B2", "B7", "B10:B22", "F5", "H10:H22", "M5"],
     "text": "Spelling: Родилеь, броллер, Прородитель, Кеегейли, Гунажу, Трик вазнда, ота на, доллрдан, Дастгох, Лойиха.",
     "resolution": "Fixed on the slides; the sheet's text is kept in the JSON as sheetText."},
]


def norm(text):
    return re.sub(r"\s+", " ", str(text)).strip()


def fix_words(text):
    for wrong, right in WORD_FIXES:
        text = text.replace(wrong, right)
    return text


def fmt(value):
    if value is None:
        return "—"
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return f"{value:,}".replace(",", " ") if isinstance(value, (int, float)) else str(value)


class Sheet:
    """One worksheet opened twice: formulas and the values Excel saved."""

    def __init__(self, wb_formulas, wb_values, name):
        self.name, self.f, self.v = name, wb_formulas[name], wb_values[name]
        self.corrected = {}
        self.notes = {}

    def ref(self, cell):
        return f"{self.name}!{cell}"

    def formula(self, cell):
        raw = self.f[cell].value
        return raw if isinstance(raw, str) and raw.startswith("=") else None

    def value(self, cell):
        return self.corrected.get(cell, self.v[cell].value)

    def num(self, cell):
        item = {"value": self.value(cell), "ref": self.ref(cell)}
        if self.formula(cell):
            item["formula"] = self.formula(cell)
        if cell in self.corrected:
            item["sheetValue"] = self.v[cell].value
            item["note"] = self.notes.get(cell, "Typed-in value in the sheet; using J+K+L+M of the same row")
        if self.ref(cell) in HIDDEN:
            item["hidden"] = True
            item["note"] = "Sum of birds per building across facilities; not shown"
        return item

    def text(self, cell, shown=None):
        raw = self.v[cell].value
        shown = fix_words(norm(raw)) if shown is None else shown
        item = {"value": shown, "ref": self.ref(cell)}
        if raw is not None and shown != raw:
            item["sheetText"] = raw
        return item


# A cell in a formula, optionally on another sheet: K8 or броллер!K8.
CELL_REF = re.compile(r"(?:(\w+)!)?([A-Z]+\d+)")


def check_cached_values(sheet, wb_values):
    """Recompute every formula from the saved inputs; a mismatch means the file wasn't recalculated."""
    problems, count = [], 0
    for row in sheet.f.iter_rows():
        for c in row:
            if not sheet.formula(c.coordinate):
                continue
            count += 1
            expr = CELL_REF.sub(lambda m: str((wb_values[m.group(1)] if m.group(1) else sheet.v)[m.group(2)].value or 0), c.value[1:])
            if not re.fullmatch(r"[\d.+\-*/() ]+", expr):
                problems.append(f"{sheet.ref(c.coordinate)}: can't recompute {c.value}")
                continue
            if abs(eval(expr) - (sheet.v[c.coordinate].value or 0)) > 1e-9:  # noqa: S307 (digits and operators only)
                problems.append(f"{sheet.ref(c.coordinate)}: saved {sheet.v[c.coordinate].value}, formula gives {eval(expr)}")
    return count, problems


def find_typed_in_totals(sheet, rows):
    """I must equal J+K+L+M on every money row. Typed-in values that don't are corrected."""
    found = []
    for base in rows:
        for offset in MONEY_ROWS.values():
            r = base + offset
            parts = sum(sheet.v[f"{c}{r}"].value or 0 for c in "JKLM")
            if not sheet.formula(f"I{r}") and (sheet.v[f"I{r}"].value or 0) != parts:
                sheet.corrected[f"I{r}"] = parts
                found.append({"ref": sheet.ref(f"I{r}"), "sheetValue": sheet.v[f"I{r}"].value, "value": parts})
    return found


def apply_adjustments(sheet):
    """Moves money between two cells as the client asked, then recalculates every formula on the sheet."""
    flags = []
    for adj in (a for a in ADJUSTMENTS if a["sheet"] == sheet.name):
        before = {c: sheet.value(c) or 0 for c in (adj["from"], adj["to"])}
        sheet.corrected[adj["from"]] = before[adj["from"]] - adj["move"]
        sheet.corrected[adj["to"]] = before[adj["to"]] + adj["move"]
        for cell in (adj["from"], adj["to"]):
            sheet.notes[cell] = adj["why"]
        for _ in range(20):  # formulas only add and multiply, so this settles in a few passes
            changed = False
            for row in sheet.f.iter_rows():
                for c in row:
                    formula = sheet.formula(c.coordinate)
                    if not formula:
                        continue
                    value = eval(re.sub(r"[A-Z]+\d+", lambda m: str(sheet.value(m.group(0)) or 0), formula[1:]))  # noqa: S307
                    if value != (sheet.value(c.coordinate) or 0):
                        sheet.corrected[c.coordinate] = value
                        sheet.notes[c.coordinate] = f"Recalculated after moving {fmt(adj['move'])} from {adj['from']} to {adj['to']}"
                        changed = True
            if not changed:
                break
        flags.append({"level": "adjustment", "refs": [sheet.ref(adj["from"]), sheet.ref(adj["to"])],
                      "text": f"{adj['why']}. {adj['from']}: {fmt(before[adj['from']])} → {fmt(sheet.value(adj['from']))}; "
                              f"{adj['to']}: {fmt(before[adj['to']])} → {fmt(sheet.value(adj['to']))}. Every total on the sheet is recalculated.",
                      "resolution": f"Section: bank {fmt(sheet.v['I9'].value)} → {fmt(sheet.value('I9'))}, own {fmt(sheet.v['I8'].value)} → {fmt(sheet.value('I8'))} thousand $."})
    return flags


def check_sums(sheet, rows):
    """After corrections every money row, facility and total must add up."""
    v, bad = sheet.value, []
    money_rows = [7, 8, 9] + [base + o for base in rows for o in MONEY_ROWS.values()]
    for r in money_rows:
        if (v(f"I{r}") or 0) != sum(v(f"{c}{r}") or 0 for c in "JKLM"):
            bad.append(f"row {r}: I ≠ J+K+L+M")
    for c in COST_COLS:
        for base in rows:
            if (v(f"{c}{base}") or 0) != (v(f"{c}{base + 1}") or 0) + (v(f"{c}{base + 2}") or 0):
                bad.append(f"{c}{base} ≠ own + bank")
        for total_row, offset in ((7, 0), (8, 1), (9, 2)):
            if (v(f"{c}{total_row}") or 0) != sum(v(f"{c}{base + offset}") or 0 for base in rows):
                bad.append(f"{c}{total_row} ≠ sum of facilities")
    for c in "CEFG":
        if (v(f"{c}7") or 0) != sum(v(f"{c}{base}") or 0 for base in rows):
            bad.append(f"{c}7 ≠ sum of facilities")
    return bad


def money(sheet, row):
    return {key: sheet.num(f"{col}{row}") for col, key in COST_COLS.items()}


# Land for each facility, in hectares. Not in the spreadsheet: the client sent it on 27.09.2026.
# Incubators got no figure. Прародитель 2-фабрика = 2 га, confirmed by the client on 27.09.2026.
LAND_SOURCE = "мижоз маълумоти, 27.09.2026"
LAND = {
    "Прородитель!B10": 3,
    "Прородитель!B16": 2,
    "Родилеь!B10": 8.4,
    "Родилеь!B13": 15,
    **{f"броллер!B{r}": 3 for r in (10, 13, 16, 19, 22)},
    **{f"Несушка!B{r}": 3 for r in (10, 13, 16)},
}
LAND_FLAG = {
    "level": "client",
    "refs": sorted(LAND),
    "text": "Land areas (ер майдони) are not in the spreadsheet; the client gave them on 27.09.2026 "
            "(Несушка 3 × 3 га, Бройлер 5 × 3 га, Прародитель 3 + 2 га, Родитель 8,4 + 15 га; incubators: none).",
    "resolution": "Shown per facility, per section and in total, with the source «мижоз маълумоти».",
}


def land_of(ref):
    return {"value": LAND[ref], "unit": "ha", "ref": LAND_SOURCE} if ref in LAND else None


def facility(sheet, section, base, warnings):
    ref = sheet.ref(f"B{base}")
    expected, shown = FACILITIES[ref]
    raw = norm(sheet.v[f"B{base}"].value)
    if raw != expected:
        warnings.append(f"{ref} text changed in the sheet; showing the sheet text. Review FACILITIES in extract.py.")
        shown = {"name": fix_words(raw)}
    supplier_raw = norm(sheet.v[f"H{base}"].value)
    if supplier_raw not in SUPPLIERS:
        warnings.append(f"{sheet.ref(f'H{base}')} supplier «{supplier_raw}» is new; showing it as written.")
    per_building, total = section["birds"]
    return {
        **shown,
        "nameRef": sheet.text(f"B{base}", shown["name"]),
        "buildings": sheet.num(f"C{base}"),
        per_building: sheet.num(f"D{base}"),
        total: sheet.num(f"E{base}"),
        "output": sheet.num(f"F{base}"),
        "revenue": sheet.num(f"G{base}"),
        "supplier": sheet.text(f"H{base}", SUPPLIERS.get(supplier_raw, supplier_raw)),
        **({"land": land_of(ref)} if ref in LAND else {}),
        **{key: money(sheet, base + offset) for key, offset in MONEY_ROWS.items()},
    }


def section_json(sheet, section, warnings):
    per_building, total = section["birds"]
    headers = {}
    for cell in HEADER_CELLS:
        if sheet.v[cell].value is not None:
            shown = HEADER_TEXT.get(sheet.ref(cell), HEADER_TEXT.get(cell))
            headers[cell] = sheet.text(cell, shown)
    title_raw = sheet.v["B2"].value
    company = re.search(r'"([^"]+)"', title_raw).group(1)
    facilities = [facility(sheet, section, base, warnings) for base in section["rows"]]
    with_land = [f for f in facilities if f.get("land")]
    return {
        "id": section["id"],
        "title": section["title"],
        "subtitle": section["subtitle"],
        "sheet": section["sheet"],
        "company": company,
        "sheetTitle": {"value": norm(title_raw), "ref": sheet.ref("B2")},
        "units": section["units"],
        "headers": headers,
        "totals": {
            "label": sheet.text("B7"),
            "buildings": sheet.num("C7"),
            per_building: sheet.num("D7"),
            total: sheet.num("E7"),
            "output": sheet.num("F7"),
            "revenue": sheet.num("G7"),
            "country": sheet.text("H7"),
            **({"land": {"value": round(sum(f["land"]["value"] for f in with_land), 2), "unit": "ha",
                         "sumOf": [f["nameRef"]["ref"] for f in with_land], "ref": LAND_SOURCE}} if with_land else {}),
            **{key: money(sheet, 7 + offset) for key, offset in MONEY_ROWS.items()},
        },
        "facilities": facilities,
    }


def summary_json(sections):
    def total(path):
        items = [s["totals"][path[0]][path[1]] if len(path) == 2 else s["totals"][path[0]] for s in sections]
        return {"value": sum(i["value"] or 0 for i in items), "sumOf": [i["ref"] for i in items]}

    hatcheries = [f["nameRef"]["ref"] for s in sections for f in s["facilities"] if f.get("kind") == "hatchery"]
    return {
        "note": "Computed from the four section sheets only. Revenue is listed, not summed: part of one section's output is another section's input.",
        **{key: {col: total((key, col)) for col in COST_COLS.values()} for key in MONEY_ROWS},
        "buildings": total(("buildings",)),
        "hatcheries": {"value": len(hatcheries), "refs": hatcheries},
        "land": {"value": round(sum(s["totals"].get("land", {}).get("value", 0) for s in sections), 2), "unit": "ha",
                 "sumOf": [ref for s in sections for ref in s["totals"].get("land", {}).get("sumOf", [])], "ref": LAND_SOURCE},
        "outputs": [{"section": s["id"], **s["totals"]["output"], "unit": s["units"]["output"]["unit"]} for s in sections],
        "revenues": [{"section": s["id"], **s["totals"]["revenue"], "unit": s["units"]["revenue"]["unit"]} for s in sections],
    }


def project_json(sheet, sections):
    """Every line and the grand totals of the project summary, checked against the section sheets.

    Returns the JSON block and the list of checks that failed.
    """
    labels = {r: norm(sheet.v[f"B{r}"].value) for r in range(1, sheet.f.max_row + 1) if sheet.v[f"B{r}"].value is not None}

    def row_of(text, prefix=True):
        rows = [r for r, t in labels.items() if (t.startswith(text) if prefix else t == text)]
        if len(rows) != 1:
            sys.exit(f"{PROJECT_SHEET}: expected one row «{text}» in column B, found {len(rows)}. The layout changed; update extract.py.")
        return rows[0]

    def block(r):
        """A line and the two rows under it: own funds, then bank credit."""
        for offset, expected in ((1, "ўз маблағи"), (2, "банк кредити")):
            if labels.get(r + offset) != expected:
                sys.exit(f"{sheet.ref(f'B{r + offset}')} is «{labels.get(r + offset)}», expected «{expected}». The layout changed; update extract.py.")
        return {key: {name: sheet.num(f"{col}{r + offset}") for col, name in PROJECT_COLS.items()} for key, offset in MONEY_ROWS.items()}

    by_id = {s["id"]: s for s in sections}
    lines = []
    for text, line_id, title in PROJECT_LINES:
        r = row_of(text)
        line = {"id": line_id, "title": title, "section": line_id in by_id, "label": sheet.text(f"B{r}", title), **block(r)}
        count = re.search(r"(\d+)\s*дона", labels[r])
        if count:
            line["count"] = {"value": int(count.group(1)), "unit": "дона", "ref": sheet.ref(f"B{r}")}
        lines.append(line)
    totals = block(row_of("Жами лойиҳалар бўйича", prefix=False))
    credit_row = row_of("Жами кредит", prefix=False)
    credit = {**sheet.num(f"D{credit_row}"), "label": sheet.text(f"B{credit_row}")["value"]}

    def v(item):
        return item["value"] or 0

    parts = [c for c in COST_COLS.values() if c != "total"]
    bad = []
    for owner in lines + [totals]:
        for key in MONEY_ROWS:
            if abs(v(owner[key]["total"]) - sum(v(owner[key][c]) for c in parts)) > 1e-6:
                bad.append(f"{owner[key]['total']['ref']} ≠ E+F+G+H")
        for col in COST_COLS.values():
            if abs(v(owner["cost"][col]) - v(owner["own"][col]) - v(owner["bank"][col])) > 1e-6:
                bad.append(f"{owner['cost'][col]['ref']} ≠ own + bank")
    for key in MONEY_ROWS:
        for col in COST_COLS.values():
            if abs(v(totals[key][col]) - sum(v(line[key][col]) for line in lines)) > 1e-6:
                bad.append(f"{totals[key][col]['ref']} ≠ sum of the lines")
    for line in (x for x in lines if x["section"]):
        section = by_id[line["id"]]["totals"]
        for key in MONEY_ROWS:
            for col in COST_COLS.values():
                if abs(v(line[key][col]) - v(section[key][col])) > 1e-6:
                    bad.append(f"{line[key][col]['ref']} = {fmt(v(line[key][col]))}, but {section[key][col]['ref']} = {fmt(v(section[key][col]))}")
    if abs(v(credit) - v(totals["bank"]["total"])) > 1e-6:
        bad.append(f"{credit['ref']} ≠ {totals['bank']['total']['ref']}")
    return {"sheet": PROJECT_SHEET, "lines": lines, **totals, "totalCredit": credit}, bad


# ---------- Қайта ишлаш (the дастгох sheet) ----------

PROCESSING_SHEET = "дастгох"
PROCESSING_GROUPS = {"slaughter": "Сўйиш, совитиш ва қадоқлаш", "cold": "Музлаткичлар", "feedmill": "Ем завод", "transport": "Махсус транспорт"}
# Each item: its row (own funds on the next row, bank credit on the one after), group, name on the
# slides and capacity. Column D («Бир соатлик қуввати бош сонда») says 3 000 on every row, which only
# fits the slaughter and chilling lines; the other capacities come from the item names (column B).
PROCESSING_ITEMS = [
    (10, "slaughter", "Сўйиш цехи", ("D", "бош/соат")),
    (19, "slaughter", "Товуқ гўштини ҳаво линиясида совитиш дастгоҳи", ("D", "бош/соат")),
    (25, "slaughter", "Қадоқлаш ва қайта ишлаш дастгоҳлари (қийма, саралаш тарозиси, қадоқлаш, вакуум)", None),
    (13, "cold", "Музлаткич (сақлаш учун)", ("B", "1 000 т")),
    (16, "cold", "Шок музлаткич", ("B", "30 т")),
    (22, "cold", "Спирал музлаткич", ("B", "3 т/соат")),
    (28, "feedmill", "Ем завод (бройлер ва тухум йўналиши учун)", ("B", "20 т/соат")),
    (34, "transport", "Озуқа ташиш учун махсус транспорт (5 × 25 т, 1 × 15 т)", ("B", "25 т ва 15 т")),
]
PROCESSING_COLS = {"F": "total", "G": "construction", "H": "equipment"}
PROCESSING_FLAGS = [
    {"level": "error", "refs": ["дастгох!F8", "дастгох!F9"],
     "text": "Own and bank totals skip rows: F8 = 400 (misses Ем завод's 300), F9 = 800 (misses every equipment row).",
     "resolution": "Using the sum of the items: own 700, bank 4 957,5. G8, G9, H8, H9 and F7 are right."},
    {"level": "total", "refs": ["дастгох!D13", "дастгох!D16", "дастгох!D22", "дастгох!D25", "дастгох!D28", "дастгох!D7"],
     "text": "«Бир соатлик қуввати бош сонда» is 3 000 on every row, also for the freezers, packaging and the feed mill; D7 adds them up to 21 000.",
     "resolution": "3 000 бош/соат shown only for the slaughter and chilling lines; other capacities from the item names."},
    {"level": "total", "refs": ["дастгох!C34", "дастгох!B34"],
     "text": "The feed trucks have no count in C. B34 lists them: бройлерга 2, несушкага 1, родитель 1 (15 т), прородительга 2.",
     "resolution": "Shown as 6 та (2 + 1 + 1 + 2), with B34 as the source."},
    {"level": "minor", "refs": ["дастгох!B37", "дастгох!A31:A45"],
     "text": "«Инкубатория (прородитель учун)» and several numbered rows have no money.",
     "resolution": "Skipped. The incubator is on the Прородитель sheet (1 380)."},
    {"level": "minor", "refs": ["Жами Лойиҳа Кегели!B24"],
     "text": "The generator (22 дона, 382,2) is a separate line of the summary, not on the дастгох sheet.",
     "resolution": "Not in the Қайта ишлаш total; the Жами slide shows it as its own line."},
]


def processing_json(sheet, project):
    """The дастгох sheet item by item, with its totals recomputed from the items."""
    v = sheet.value
    for cell, expected in (("B7", "Қайта ишлаш йўналиши бўйича"), ("B8", "ўз маблағи"), ("B9", "банк кредити")):
        if norm(v(cell)) != expected:
            sys.exit(f"{sheet.ref(cell)} is «{v(cell)}», expected «{expected}». The layout changed; update extract.py.")
    items, bad = [], []
    for row, group, name, capacity in PROCESSING_ITEMS:
        for offset, expected in ((1, "ўз маблағи"), (2, "банк кредити")):
            if norm(v(f"B{row + offset}")) != expected:
                sys.exit(f"{sheet.ref(f'B{row + offset}')} is «{v(f'B{row + offset}')}», expected «{expected}». The layout changed; update extract.py.")
        money = {key: {name_: sheet.num(f"{col}{row + off}") for col, name_ in PROCESSING_COLS.items()}
                 for key, off in (("cost", 0), ("own", 1), ("bank", 2))}
        # own/bank F cells are sometimes empty: the row's total is G + H
        for key in ("own", "bank"):
            f = money[key]["total"]
            parts = (money[key]["construction"]["value"] or 0) + (money[key]["equipment"]["value"] or 0)
            if (f["value"] or 0) != parts:
                f.update({"sheetValue": f["value"], "value": parts, "note": "Empty or different in the sheet; G + H of the same row"})
        if sheet.v[f"C{row}"].value is not None:
            count = sheet.num(f"C{row}")
        else:
            listed = [int(n) for n in re.findall(r"(\d+)\s*дона", norm(v(f"B{row}")))]
            count = {"value": sum(listed), "ref": sheet.ref(f"B{row}"), "note": " + ".join(map(str, listed)) + " дона in the name; C is empty"}
        cap = None
        if capacity and capacity[0] == "D":
            cap = {"value": f"{fmt(v(f'D{row}'))} {capacity[1]}", "ref": sheet.ref(f"D{row}")}
        elif capacity:
            cap = {"value": capacity[1], "ref": sheet.ref(f"B{row}")}
        items.append({"row": row, "group": group, "name": sheet.text(f"B{row}", name), "count": count,
                      "capacity": cap, "country": sheet.text(f"E{row}"), **money})

    def n(item):
        return item["value"] or 0

    for it in items:
        for col in PROCESSING_COLS.values():
            if abs(n(it["cost"][col]) - n(it["own"][col]) - n(it["bank"][col])) > 1e-6:
                bad.append(f"{it['cost'][col]['ref']} ≠ own + bank")
        for key in ("cost", "own", "bank"):
            if abs(n(it[key]["total"]) - n(it[key]["construction"]) - n(it[key]["equipment"])) > 1e-6:
                bad.append(f"{it[key]['total']['ref']} ≠ G + H")
    totals = {}
    for key, row in (("cost", 7), ("own", 8), ("bank", 9)):
        totals[key] = {}
        for col, name_ in PROCESSING_COLS.items():
            computed = sum(n(it[key][name_]) for it in items)
            item = sheet.num(f"{col}{row}")
            if abs((item["value"] or 0) - computed) > 1e-6:
                item.update({"sheetValue": item["value"], "value": computed, "note": "Formula skips rows; using the sum of the items"})
            totals[key][name_] = item
    line = next(l for l in project["lines"] if l["id"] == "processing")
    for key in ("cost", "own", "bank"):
        diff = n(totals[key]["total"]) - n(line[key]["total"])
        if abs(diff) > 0.5 + 1e-6:
            bad.append(f"{totals[key]['total']['ref']} ({fmt(n(totals[key]['total']))}) ≠ {line[key]['total']['ref']} ({fmt(n(line[key]['total']))})")
    return {"sheet": PROCESSING_SHEET, "groups": PROCESSING_GROUPS, "items": items, **totals}, bad


# ---------- verification.md ----------

def cell_md(item):
    if item is None or item.get("value") is None:
        return "—"
    text = fmt(item["value"])
    if "sheetValue" in item:
        return f"**{text}** ⚠ (sheet: {fmt(item['sheetValue'])})"
    if item.get("hidden"):
        return f"~~{text}~~"
    return f"*{text}*" if "formula" in item else text


def matrix_md(sec, section):
    per_building, total = section["birds"]
    lab = {"C": "Бино", "I": "Жами", "J": "Қурилиш", "K": "Дастгоҳ", "M": "Озуқа", **section["labels"]}
    cols = "CDEFGIJKLM"
    lines = ["| Row | B | " + " | ".join(f"{c} · {lab[c]}" for c in cols) + " |",
             "|---:|---|" + "---:|" * len(cols)]

    def row(r, label, item_for):
        lines.append(f"| {r} | {label} | " + " | ".join(cell_md(item_for(c)) for c in cols) + " |")

    def top(t):
        keyed = {"C": t.get("buildings"), "D": t.get(per_building), "E": t.get(total),
                 "F": t.get("output"), "G": t.get("revenue")}
        return lambda c: keyed[c] if c in keyed else t["cost"][COST_COLS[c]]

    def sub(t, key):
        return lambda c: t[key][COST_COLS[c]] if c in COST_COLS else None

    t = sec["totals"]
    row(7, "**" + t["label"]["value"] + "**", top(t))
    row(8, "↳ ўз маблағи", sub(t, "own"))
    row(9, "↳ банк кредити", sub(t, "bank"))
    for base, f in zip(section["rows"], sec["facilities"]):
        row(base, "**" + f["name"] + "**", top(f))
        row(base + 1, "↳ ўз маблағи", sub(f, "own"))
        row(base + 2, "↳ банк кредити", sub(f, "bank"))
    return lines


def text_md(sec):
    lines = ["| Cell | In the sheet | On the slides |", "|---|---|---|"]
    items = [sec["sheetTitle"] | {"shown": f"{sec['title']} — {sec['subtitle']}"}, sec["totals"]["label"], sec["totals"]["country"]]
    for f in sec["facilities"]:
        extra = " · ".join(x for x in (f.get("size"), f.get("capacity"), f.get("district")) if x)
        items.append(f["nameRef"] | {"shown": f["name"] + (f" · {extra}" if extra else "")})
        items.append(f["supplier"])
    for it in items:
        sheet_text = norm(it.get("sheetText", it["value"]))
        lines.append(f"| {it['ref'].split('!')[1]} | {sheet_text} | {it.get('shown', it['value'])} |")
    return lines


def project_md(project):
    """The project summary laid out like the sheet: each line, its own funds and its bank credit."""
    names = {"D": "Жами", "E": "Қурилишга", "F": "Дастгоҳга", "G": "Жўжа", "H": "Озуқа"}
    lines = ["| Row | B | " + " | ".join(f"{c} · {names[c]}" for c in PROJECT_COLS) + " |", "|---:|---|" + "---:|" * len(PROJECT_COLS)]

    def add(owner, label):
        for key, row_label in (("cost", f"**{label}**"), ("own", "↳ ўз маблағи"), ("bank", "↳ банк кредити")):
            row = owner[key]
            lines.append(f"| {row['total']['ref'].split('!')[1][1:]} | {row_label} | "
                         + " | ".join(cell_md(row[name]) for name in PROJECT_COLS.values()) + " |")

    for line in project["lines"]:
        count = f" · {line['count']['value']} {line['count']['unit']}" if line.get("count") else ""
        add(line, f"{line['title']}{count}")
    add(project, "Жами лойиҳалар бўйича")
    credit = project["totalCredit"]
    lines.append(f"| {credit['ref'].split('!')[1][1:]} | **{credit['label']}** | {cell_md(credit)} | | | | |")
    return lines


def full_list_md(sections):
    lines = ["| # | Section | What | Value | Cell | Formula | Note |", "|---:|---|---|---:|---|---|---|"]
    n = 0
    names = {"total": "Лойиҳа қиймати", "construction": "Қурилишга", "equipment": "Дастгоҳга", "chickens": "Товуқ/жўжа", "feed": "Озуқа"}
    parts = {"cost": "", "own": " · ўз маблағи", "bank": " · банк кредити"}

    def add(section, what, item):
        nonlocal n
        n += 1
        formula = f"`{item['formula']}`" if item.get("formula") else ""
        note = item.get("note", "")
        if "sheetValue" in item:
            note = f"sheet has {fmt(item['sheetValue'])}; {note}"
        lines.append(f"| {n} | {section} | {what} | {fmt(item['value'])} | {item['ref']} | {formula} | {note} |")

    bird_names = {"birdsPerBuilding": "бош (1 бинода)", "birds": "жами бош",
                  "birdsPerBatch": "бош (1 боқишда)", "birdsPerYear": "бош (йилига)"}
    for s in sections:
        pb, tb = s["birds_keys"]
        for owner, prefix in [(s["totals"], "Жами")] + [(f, f["name"]) for f in s["facilities"]]:
            add(s["title"], f"{prefix} · бино сони", owner["buildings"])
            add(s["title"], f"{prefix} · {bird_names[pb]}", owner[pb])
            add(s["title"], f"{prefix} · {bird_names[tb]}", owner[tb])
            add(s["title"], f"{prefix} · маҳсулот", owner["output"])
            add(s["title"], f"{prefix} · тушум", owner["revenue"])
            for key, suffix in parts.items():
                for col, name in names.items():
                    add(s["title"], f"{prefix}{suffix} · {name}", owner[key][col])
    return lines


def write_markdown(data, checks, flags, merged, warnings):
    order = {s["id"]: s for s in SECTIONS}
    md = [
        "# Кегейли.xlsx → kegeyli.json: check table",
        "",
        f"Source `{data['meta']['source']}` · SHA-256 `{data['meta']['sha256'][:16]}…` · extracted {data['meta']['extractedAt']}",
        "",
        "Regenerate with `.venv/bin/python scripts/extract.py`. Every number below comes from the JSON the slides use.",
        "",
        "**How to read the tables:** the cell address is the column letter plus the row number (column `I` in row `7` is `I7`).",
        "*Italic* = the cell holds a formula in Excel · **bold ⚠** = changed from the sheet · ~~struck~~ = not shown on slides · — = empty cell.",
        "",
        "## Automatic checks",
        "",
        *[f"- {line}" for line in checks],
        *([f"- ⚠ {w}" for w in warnings] or []),
        "",
        "## Flags",
        "",
        "| # | Kind | Cells | Problem | What we do |",
        "|---:|---|---|---|---|",
        *[f"| {i} | {f['level']} | {', '.join(f['refs'])} | {f['text']} | {f['resolution']} |" for i, f in enumerate(flags, 1)],
        "",
        "## Merged cells",
        "",
        "Only titles and column headers are merged; no number sits inside a merged range. The text lives in the top-left cell.",
        "",
        "| Sheet | Merged ranges |",
        "|---|---|",
        *[f"| {name} | {', '.join(ranges)} |" for name, ranges in merged.items()],
        "",
    ]
    for i, sec in enumerate(data["sections"], 2):
        section = order[sec["id"]]
        u = sec["units"]
        md += [
            f"## Slide {i} · {sec['title']} — sheet `{sec['sheet']}`",
            "",
            f"Units: I–M in **{u['cost']['unit']}** · G revenue in **{u['revenue']['unit']}** · F output in **{u['output']['unit']}** ({u['output']['what']})",
            "",
            *text_md(sec),
            "",
            *matrix_md(sec, section),
            "",
        ]
    s = data["summary"]
    md += [
        f"## Slides 1 and 6 · Жами — sheet `{PROJECT_SHEET}`",
        "",
        "The cover and the Жами slide use these totals. The four section rows are checked against their own sheets (see Automatic checks).",
        "",
        *project_md(data["project"]),
        "",
        "## The four sections added up (buildings and land on the Жами slide)",
        "",
        "| What | Value, thousand $ | Sum of |",
        "|---|---:|---|",
        *[f"| {part} · {col} | {fmt(s[part][col]['value'])} | {', '.join(s[part][col]['sumOf'])} |"
          for part in MONEY_ROWS for col in COST_COLS.values()],
        f"| buildings | {fmt(s['buildings']['value'])} | {', '.join(s['buildings']['sumOf'])} |",
        f"| incubators among them | {s['hatcheries']['value']} | {', '.join(s['hatcheries']['refs'])} |",
        "",
        f"## Ер майдони ({LAND_SOURCE}, not in the spreadsheet)",
        "",
        "| Section | Facility | Land, га |",
        "|---|---|---:|",
        *[f"| {sec['title']} | {f['name']} | {fmt(f['land']['value']) if f.get('land') else '—'} |"
          for sec in data["sections"] for f in sec["facilities"]],
        *[f"| **{sec['title']}** | **жами** | **{fmt(sec['totals']['land']['value'])}** |" for sec in data["sections"] if sec["totals"].get("land")],
        f"| **Жами** | | **{fmt(s['land']['value'])}** |",
        "",
        "## Full list (one line per value)",
        "",
        *full_list_md([{**sec, "birds_keys": order[sec["id"]]["birds"]} for sec in data["sections"]]),
        "",
    ]
    OUT_MD.parent.mkdir(parents=True, exist_ok=True)
    OUT_MD.write_text("\n".join(md), encoding="utf-8")


def main():
    if not SOURCE.exists():
        sys.exit(f"Spreadsheet not found: {SOURCE}")
    wb_f = openpyxl.load_workbook(SOURCE, data_only=False)
    wb_v = openpyxl.load_workbook(SOURCE, data_only=True)

    checks, flags, warnings, merged, sections = [], [], [], {}, []
    for section in SECTIONS:
        sheet = Sheet(wb_f, wb_v, section["sheet"])
        for cell, expected in (("B8", "ўз маблағи"), ("B9", "банк кредити"), ("I4", "Лойиҳани қиймати")):
            if norm(sheet.v[cell].value) != expected:
                sys.exit(f"{sheet.ref(cell)} is «{sheet.v[cell].value}», expected «{expected}». The layout changed; update extract.py.")
        count, stale = check_cached_values(sheet, wb_v)
        if stale:
            sys.exit("Saved values don't match the formulas. Open the file in Excel, save it, and run again:\n" + "\n".join(stale))
        checks.append(f"`{sheet.name}`: all {count} formulas give the value Excel saved.")
        typed_fixes = find_typed_in_totals(sheet, section["rows"])
        for fix in typed_fixes:
            flags.append({"level": "error", "refs": [fix["ref"]],
                          "text": f"Typed-in {fmt(fix['sheetValue'])} where J+K+L+M of the same row gives {fmt(fix['value'])}; the two sub-rows then don't add up to the facility row.",
                          "resolution": f"Using {fmt(fix['value'])}. Section totals (rows 7–9) are built from J–M, so they don't change."})
        adjustments = apply_adjustments(sheet)
        flags.extend(adjustments)
        bad = check_sums(sheet, section["rows"])
        done = [f"the fix in {', '.join(f['ref'].split('!')[1] for f in typed_fixes)}"] if typed_fixes else []
        done += ["the client's adjustment"] if adjustments else []
        after = f" (after {' and '.join(done)})" if done else ""
        checks.append(f"`{sheet.name}`: own + bank = facility, facilities = total, I = J+K+L+M on every row{after}: "
                      + ("all pass." if not bad else "FAIL — " + "; ".join(bad)))
        merged[sheet.name] = [f"{r} («{norm(sheet.v[str(r).split(':')[0]].value or '')[:40]}»)"
                              for r in sorted(sheet.f.merged_cells.ranges, key=str)]
        sections.append(section_json(sheet, section, warnings))

    project_sheet = Sheet(wb_f, wb_v, PROJECT_SHEET)
    count, stale = check_cached_values(project_sheet, wb_v)
    if stale:
        sys.exit("Saved values don't match the formulas. Open the file in Excel, save it, and run again:\n" + "\n".join(stale))
    checks.append(f"`{PROJECT_SHEET}`: all {count} formulas give the value Excel saved.")
    project, bad = project_json(project_sheet, sections)
    checks.append(f"`{PROJECT_SHEET}`: section rows = their sheets, own + bank = line, D = E+F+G+H, lines add up to Жами, "
                  "Жами кредит = банк кредити: " + ("all pass." if not bad else "FAIL — " + "; ".join(bad)))
    merged[PROJECT_SHEET] = [f"{r} («{norm(project_sheet.v[str(r).split(':')[0]].value or '')[:40]}»)"
                             for r in sorted(project_sheet.f.merged_cells.ranges, key=str)]

    processing_sheet = Sheet(wb_f, wb_v, PROCESSING_SHEET)
    count, stale = check_cached_values(processing_sheet, wb_v)
    if stale:
        sys.exit("Saved values don't match the formulas. Open the file in Excel, save it, and run again:\n" + "\n".join(stale))
    checks.append(f"`{PROCESSING_SHEET}`: all {count} formulas give the value Excel saved.")
    processing, bad = processing_json(processing_sheet, project)
    checks.append(f"`{PROCESSING_SHEET}`: own + bank = item, F = G + H, items add up to rows 7–9 (after fixing F8, F9), "
                  "total = the summary's Қайта ишлаш line (±0,5): " + ("all pass." if not bad else "FAIL — " + "; ".join(bad)))

    companies = {s["company"] for s in sections}
    if len(companies) != 1:
        warnings.append(f"Company name differs between sheets: {companies}")
    data = {
        "meta": {
            "source": str(SOURCE.relative_to(ROOT)) if SOURCE.is_relative_to(ROOT) else str(SOURCE),
            "sha256": hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
            "extractedAt": dt.datetime.now().isoformat(timespec="seconds"),
            "company": sections[0]["company"],
            "units": {"kUSD": "thousand US dollars", "USD": "US dollars", "kg": "kilograms", "pcs": "pieces"},
        },
        "sections": sections,
        "summary": summary_json(sections),
        "project": project,
        "processing": processing,
        "flags": flags + STATIC_FLAGS + PROJECT_FLAGS + PROCESSING_FLAGS + [LAND_FLAG],
    }
    OUT_JSON.parent.mkdir(parents=True, exist_ok=True)
    OUT_JSON.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    write_markdown(data, checks, data["flags"], merged, warnings)
    print(f"Wrote {OUT_JSON.relative_to(ROOT)} and {OUT_MD.relative_to(ROOT)}")
    for line in checks + [f"WARNING: {w}" for w in warnings]:
        print(" -", line)


if __name__ == "__main__":
    main()
